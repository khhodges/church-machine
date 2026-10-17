"""Append the bounded audit to the workbook without resetting existing rows."""
from copy import copy
import importlib.util
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "docs/phase-1/CTMM_W7_Threat_Test_Plan.xlsx"
REPORT = "docs/reviews/phase1-baseline/README.md"


def values(ws):
    return [(c.coordinate, c.value, str(c._style)) for row in ws for c in row if c.value is not None]


def main():
    book = openpyxl.load_workbook(BOOK)
    assert "Baseline audit" not in book.sheetnames, "Audit already recorded; do not reset user edits."
    original = {n: values(book[n]) for n in ["Threat Map", "Summary", "Method"]}
    tracker = book["Completeness"]
    prior = {tracker.cell(r, 1).value: tracker.cell(r, 6).value for r in range(2, tracker.max_row+1)}
    audit = book.create_sheet("Baseline audit")
    audit.append(["ID", "Area", "Classification / confidence", "Observation", "Next action", "Related requirements"])
    items = []
    for line in (ROOT / REPORT).read_text().splitlines():
        if line.startswith("| BASE-"):
            item = [s.strip().replace("`", "") for s in line.strip("|").split("|")]
            assert len(item) == 6
            items.append(item)
            audit.append(item)
    assert len(items) == 9
    for id, area, classification, observation, action, related in items:
        assert id not in prior
        tracker.append([id, "Baseline audit", area, "Phase 1 assurance", "High",
                        "Needs verification" if id in ("BASE-05", "BASE-08") else "Open",
                        "Unassigned", related, classification + ". " + observation,
                        action, REPORT, "2026-10-17", REPORT, None])
        r = tracker.max_row
        for col in range(1, 15):
            tracker.cell(r, col)._style = copy(tracker.cell(2, col)._style)
        tracker.row_dimensions[r].height = 100
        for row in range(2, tracker.max_row):
            if tracker.cell(row, 1).value in related.split("; "):
                cell = tracker.cell(row, 13)
                cell.value = (str(cell.value) + "; " if cell.value else "") + id + ": " + REPORT
    tracker.tables["Phase1Tracker"].ref = f"A1:N{tracker.max_row}"
    tracker.auto_filter.ref = f"A1:N{tracker.max_row}"
    widths = [16, 26, 55, 95, 90, 55]
    for i, width in enumerate(widths, 1):
        audit.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    for c in audit[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17365D")
    for row in audit.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        audit.row_dimensions[row[0].row].height = 110
    table = Table(displayName="BaselineAudit", ref=audit.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    audit.add_table(table)
    audit.freeze_panes = "C2"
    audit.sheet_view.showGridLines = False
    audit.print_title_rows = "1:1"
    audit.page_setup.orientation = "landscape"
    audit.page_setup.fitToWidth = 1
    audit.sheet_properties.pageSetUpPr.fitToPage = True
    # Add the new area to the dynamic summary rather than excluding it silently.
    progress = book["Progress"]
    n = progress.max_row + 1
    progress.append(["Baseline audit", f'=COUNTIF(Completeness!B2:B10000,A{n})',
                     f'=COUNTIFS(Completeness!B2:B10000,A{n},Completeness!F2:F10000,"Complete")',
                     f'=COUNTIFS(Completeness!B2:B10000,A{n},Completeness!F2:F10000,"Blocked")'])
    for col in range(1, 5):
        progress.cell(n, col)._style = copy(progress.cell(n-1, col)._style)
    book["Tracking guide"].append(["Baseline audit", "Nine root-cause groups. Workflow results are existing observations, not fresh reruns. Only the compiler dependency and read-only hash/binding probes were reproduced. See repository audit report."])
    candidate = BOOK.with_suffix(".candidate.xlsx")
    book.save(candidate)
    check = openpyxl.load_workbook(candidate)
    assert all(values(check[n]) == data for n, data in original.items())
    current = {r[0].value: r[5].value for r in list(check["Completeness"])[1:]}
    assert all(current[id] == status for id, status in prior.items())
    assert len(current) == len(prior) + 9
    assert check["Baseline audit"].max_row == 10
    candidate.replace(BOOK)
    print(f"Validated {len(current)} tracker rows; original sheets and {len(prior)} prior statuses preserved.")


if __name__ == "__main__":
    main()
