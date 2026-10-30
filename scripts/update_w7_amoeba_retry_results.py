"""Record bounded retry acceptance without rewriting historical workbook parts."""
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import posixpath
import xml.etree.ElementTree as ET
import openpyxl
from refresh_phase1_progress import refresh, NS, REL

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "docs/phase-1/CTMM_W7_Threat_Test_Plan.xlsx"
DATE = "2026-10-30"
REPORT = "reports/amoeba-retry-candidate.md"
BUILD = (
    "Amoeba retry candidate build PASS: Vivado 2026.1, exit 0; setup +2.927 ns, "
    "hold +0.039 ns, pulse width +7.000 ns; zero reported timing failures. "
    "Uploaded vendor archive independently verified: 126 input hashes and 7 output hashes; "
    "manifest matches the locally prepared package. "
    "BIT SHA256 a4acb0ef9a39ba985d7769031cf146edb68fb65ff96e0b4f4a895093b54e5d4e; "
    "image SHA256 5603d8e105e47ca0f6524e4553989387c088de19ac0e6b3f6fd9897e170cf16d. "
    "No configuration-flash write or general production-release approval.")
PHYSICAL = (
    "Bounded physical retry PASS: operator checksum-verified the exact BIT and supplied "
    "successful temporary-programming output for the recorded Wukong board. Startup plus "
    "six explicitly commanded manual reboots yielded seven complete CRC-valid fault "
    "snapshots, sequences 1 through 7. Each fault: code 8 at 0x3A08, Mallory CR12 "
    "0x0200000C. Every retry traversed 0x4638, not old root RETURN 0x463C. "
    "Alice advanced 42→43→44→45→46→47→48; no Boot/Alice fault or STACK_UNDERFLOW observed. "
    "Snapshot register/stack fields unchanged except sequence/CRC. This is stable exposed "
    "state, not an exhaustive stack-memory audit, indefinite recovery, automatic fault "
    "release, or live installed-image readback attestation.")
RTL = (
    "Exact-image frozen Amaranth RTL replay PASS: first intended Mallory fault, simulated "
    "manual reboot f, 1,854 post-reboot retirements across Boot/Alice/Mallory, retry branch, "
    "then the same intended Mallory fault. No root underflow. Frozen checksums unchanged. "
    "This is not a full vendor-mapped recovery replay.")
UPDATES = {
    "REL-01": ("Complete", BUILD,
               "Exact candidate build-provenance criterion satisfied. Keep immutable history; broader release/physical gates remain separate."),
    "BASE-09": ("Complete", BUILD,
                "Exact candidate provenance gap closed; retain historical source-drift incidents without relabeling historical BITs."),
    "REL-02": ("Needs verification", PHYSICAL + " " + BUILD,
               "Bounded physical startup/retry passed. Complete remaining H19/H20 objectives before closing full physical acceptance; no flash approval."),
    "H20": (None, PHYSICAL,
            "Retain six successful manual retries; complete separately authorized cold/power-cycle and remaining retained-control/initialization objectives at the original scope."),
    "H19": (None, RTL + " " + PHYSICAL,
            "Complete same-fixture behavioral/emitted/vendor-mapped/board committed-state comparisons. UART traces and fault snapshots are bounded observations, not complete memory equality."),
    "CM-H": (None, BUILD + " " + PHYSICAL,
             "Exact BIT/image/build evidence is preserved for this candidate. Complete item-specific identity coverage and distinguish operator programming provenance from live data-image readback."),
    "SIM-03": (None, RTL,
               "Do not infer full CHANGE/SWITCH, Namespace geometry or all-profile conformance from this retry fixture."),
    "SIM-04": (None, RTL + " " + PHYSICAL,
               "Keep remaining watchdog, competing-fault precedence and elapsed-time objectives open; manual f is not automatic fault authorization."),
    "BASE-05": (None, RTL + " " + PHYSICAL,
                "Retain unrelated historical regression failures until individually rerun. This test does not close the whole RTL suite."),
    "BASE-06": (None, BUILD + " " + PHYSICAL,
                "This candidate regenerated and booted physically. Retain broader generator/initializer and cold-start objectives until separately qualified."),
}


def main():
    before = openpyxl.load_workbook(BOOK)
    before_cached = openpyxl.load_workbook(BOOK, data_only=True)
    indexes = {r[0].value: r[0].row for r in before["Completeness"] if r[0].value}
    with ZipFile(BOOK) as z:
        infos = z.infolist()
        original = {i.filename: z.read(i.filename) for i in infos}
    files = dict(original)
    wb = ET.fromstring(files["xl/workbook.xml"])
    rels = ET.fromstring(files["xl/_rels/workbook.xml.rels"])
    targets = {r.get("Id"): r.get("Target") for r in rels}
    sheets = {}
    for s in wb.find(f"{{{NS}}}sheets"):
        target = targets[s.get(f"{{{REL}}}id")]
        sheets[s.get("name")] = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
    trees = {n: ET.fromstring(files[sheets[n]])
             for n in ("Completeness", "Progress", "Tracking guide", "Plan guide")}

    def set_cell(sheet, address, value):
        c = next(c for c in trees[sheet].iter(f"{{{NS}}}c") if c.get("r") == address)
        for child in list(c):
            c.remove(child)
        c.set("t", "inlineStr")
        ET.SubElement(ET.SubElement(c, f"{{{NS}}}is"), f"{{{NS}}}t").text = value

    prefix = f"Update {DATE} — Amoeba retry: "
    action_prefix = f"Amoeba retry next action ({DATE}): "
    for key, (status, evidence, action) in UPDATES.items():
        row = indexes[key]
        if status:
            set_cell("Completeness", f"F{row}", status)
        for col, lead, text, separator in (
            ("I", prefix, evidence, "\n\nEarlier evidence: "),
            ("J", action_prefix, action, "\n\nOriginal criterion / prior actions: "),
        ):
            old = before["Completeness"][f"{col}{row}"].value or ""
            if old.startswith(lead):
                old = old.split(separator, 1)[1]
            set_cell("Completeness", f"{col}{row}", lead + text + separator + old)
        old_links = before["Completeness"][f"M{row}"].value or ""
        links = old_links.split("; ") if old_links else []
        for link in (REPORT, "generated-artifacts/amoeba-retry-preserved/vendor-verification.json",
                     "generated-artifacts/amoeba-retry-preserved/amoeba-retry-complete-evidence.zip"):
            if link not in links:
                links.append(link)
        set_cell("Completeness", f"M{row}", "; ".join(links))
        set_cell("Completeness", f"L{row}", DATE)
    set_cell("Progress", "B2", DATE)
    summary = (
        "30 October: exact retry BIT build/provenance PASS, frozen Amaranth RTL retry PASS, "
        "physical startup and six manual retries PASS (snapshots 1–7; Alice 42→48; "
        "only intended Mallory faults). REL-01 and BASE-09 provenance complete. "
        "Full H19/H20 and REL-02 acceptance remains Needs verification; no new W7 threat "
        "closure, automatic recovery, flash programming or production-release approval. "
        "Original claims, scope baseline and historical evidence preserved.")
    set_cell("Tracking guide", "B7", summary)
    set_cell("Plan guide", "B11", summary)
    for n, tree in trees.items():
        files[sheets[n]] = ET.tostring(tree, encoding="utf-8")
    candidate = BOOK.with_suffix(".retry-update.xlsx")
    with ZipFile(candidate, "w") as z:
        for info in infos:
            z.writestr(info, files[info.filename])
    refresh(candidate)
    with ZipFile(candidate) as z:
        changed = {sheets[n] for n in trees}
        assert all(z.read(n) == data for n, data in original.items() if n not in changed)
    after = openpyxl.load_workbook(candidate)
    cached = openpyxl.load_workbook(candidate, data_only=True)
    for s in before:
        for row in s:
            for c in row:
                if c.data_type == "f":
                    assert after[s.title][c.coordinate].value == c.value
                    if s.title == "Progress":
                        assert cached[s.title][c.coordinate].value is not None
                    else:
                        assert cached[s.title][c.coordinate].value == before_cached[s.title][c.coordinate].value
    for key, row in indexes.items():
        if key not in UPDATES:
            assert [c.value for c in after["Completeness"][row]] == [c.value for c in before["Completeness"][row]]
    counts = Counter(r[5] for r in list(cached["Completeness"].values)[1:])
    assert counts["Complete"] == 4 and cached["Progress"]["B3"].value == 113
    assert cached["Progress"]["B7"].value == counts["Needs verification"]
    candidate.replace(BOOK)
    print("PASS: 10 rows updated; all unrelated workbook parts and formulas preserved; cached results verified.")
    print(dict(counts))


if __name__ == "__main__":
    main()
