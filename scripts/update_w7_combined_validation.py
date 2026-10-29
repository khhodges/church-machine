"""Record bounded combined validation, preserving original workbook sheets."""
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import posixpath
import xml.etree.ElementTree as ET
import openpyxl
from refresh_phase1_progress import refresh, NS, REL

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "docs/phase-1/CTMM_W7_Threat_Test_Plan.xlsx"
DATE = "2026-10-29"
REPORT = "docs/phase-1/combined-validation-status-20261029.md"
BOOT = ("Combined frozen candidate: vendor-synthesized initialization/boot PASS; "
        "all 2,460 address/data writes checked, entry 0x4610 at 25 MHz. "
        "Original and corrected-observer checksums verified; vendor runner exit 0.")
MODEL = ("Combined recovery-model PASS: model runner exit 0, subsequently accepted "
         "by the build preflight requiring the explicit pass marker and result JSON; "
         "one fault, authorized boot recovery, 100 post-recovery retirements, third "
         "address 0x2870. This is Amaranth simulated-UART evidence, not vendor-mapped recovery.")
BUILD = ("Build IN PROGRESS as last reported: Vivado 2026.1 opened the validated "
         "checkpoint on 2026-10-29 at 18:13:31 (tool-reported time). No completed "
         "routing/timing/DRC result or combined BIT yet reported. No combined board programming.")
UPDATES = {
    "H19": ("Needs verification", BOOT + " " + MODEL,
            "Complete same-fixture cross-layer state comparison and separately authorized physical replay; boot-only vendor evidence does not establish mapped recovery."),
    "H20": ("Needs verification", MODEL + " " + BOOT,
            "Perform exact combined-build cold/warm boots, repeated handoffs and physical recovery after separate programming approval."),
    "CM-H": (None, BOOT + " " + BUILD,
             "Retain image/source/checkpoint/evidence hashes, then bind the routed BIT receipt and physical observations; do not relabel the earlier BIT."),
    "SIM-03": (None, BOOT + " " + MODEL,
               "Retain this exact-image result without claiming full CHANGE/SWITCH or all-profile conformance."),
    "SIM-04": (None, MODEL,
               "Verify remaining first-fault/precedence and physical timing objectives; model bounds are not FPGA elapsed-time measurements."),
    "BASE-05": (None, MODEL + " " + BOOT,
                "Retain earlier unrelated RTL failures as unresolved until individually rerun; this candidate does not close the whole regression suite."),
    "BASE-06": ("Needs verification", BOOT,
                "Confirm combined initializer on the physical board after the routed build passes. Historical generator/count failures are not all closed by this isolated candidate."),
    "BASE-09": ("Needs verification", BUILD,
                "Collect routed timing, DRC, output hashes and build receipt; historical source drift remains historical, not a reason to relabel old artifacts."),
    "REL-01": ("Needs verification", BOOT + " " + MODEL + " " + BUILD,
               "Review COMBINED_IMPLEMENTATION_PASS, exit 0, timing/DRC reports and hash-bearing receipt before accepting the experimental build."),
    "REL-02": ("Not run", BUILD + " Earlier continuation-only BIT failed startup; "
               "that incident is not a physical result for the combined candidate.",
               "Obtain separate programming approval, temporarily load the exact reviewed BIT, then retain physical startup/recovery evidence; persistent flash stays unchanged."),
}


def main():
    before = openpyxl.load_workbook(BOOK)
    indexes = {r[0].value: r[0].row for r in before["Completeness"] if r[0].value}
    with ZipFile(BOOK) as z:
        infos = z.infolist()
        original = {i.filename: z.read(i.filename) for i in infos}
    files = dict(original)
    wb = ET.fromstring(files["xl/workbook.xml"])
    rels = ET.fromstring(files["xl/_rels/workbook.xml.rels"])
    targets = {r.get("Id"): r.get("Target") for r in rels}
    sheets = {}
    for s in wb.find(f"{{{NS}}}sheets"):
        target = targets[s.get(f"{{{REL}}}id")]
        sheets[s.get("name")] = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
    trees = {n: ET.fromstring(files[sheets[n]]) for n in ("Completeness", "Progress", "Tracking guide")}

    def set_cell(sheet, address, value):
        cells = {c.get("r"): c for c in trees[sheet].iter(f"{{{NS}}}c")}
        c = cells[address]
        for child in list(c):
            c.remove(child)
        c.set("t", "inlineStr")
        ET.SubElement(ET.SubElement(c, f"{{{NS}}}is"), f"{{{NS}}}t").text = value

    prefix = f"Update {DATE} — combined candidate: "
    for key, (status, evidence, action) in UPDATES.items():
        row = indexes[key]
        if status:
            set_cell("Completeness", f"F{row}", status)
        for col, text, separator in (
            ("I", prefix + evidence, "\n\nEarlier evidence: "),
            ("J", f"Combined-candidate next action ({DATE}): " + action, "\n\nOriginal criterion / prior actions: "),
        ):
            old = before["Completeness"][f"{col}{row}"].value or ""
            if old.startswith(prefix if col == "I" else f"Combined-candidate next action ({DATE}): "):
                old = old.split(separator, 1)[1]
            set_cell("Completeness", f"{col}{row}", text + separator + old)
        links = before["Completeness"][f"M{row}"].value or ""
        if REPORT not in links.split("; "):
            links = REPORT + ("; " + links if links else "")
        set_cell("Completeness", f"M{row}", links)
        set_cell("Completeness", f"L{row}", DATE)
    set_cell("Progress", "B2", DATE)
    set_cell("Tracking guide", "B7",
             "29 October: user-supplied droplet logs and build preflight establish combined vendor boot and recovery-model passes. Implementation last reported in progress; combined physical acceptance not run. No W7 threat newly passed. Original threat sheets and historical criteria retained.")
    for name, tree in trees.items():
        files[sheets[name]] = ET.tostring(tree, encoding="utf-8")
    candidate = BOOK.with_suffix(".combined-update.xlsx")
    with ZipFile(candidate, "w") as z:
        for info in infos:
            z.writestr(info, files[info.filename])
    refresh(candidate)
    with ZipFile(candidate) as z:
        changed = {sheets[n] for n in trees}
        assert all(z.read(n) == data for n, data in original.items() if n not in changed)
    after = openpyxl.load_workbook(candidate)
    cached = openpyxl.load_workbook(candidate, data_only=True)
    for s in before:
        for row in s:
            for cell in row:
                if cell.data_type == "f":
                    assert after[s.title][cell.coordinate].value == cell.value
    for key, row in indexes.items():
        if key not in UPDATES:
            assert [c.value for c in after["Completeness"][row]] == [c.value for c in before["Completeness"][row]]
    counts = Counter(r[5] for r in list(cached["Completeness"].values)[1:])
    assert counts["Complete"] == 2
    assert cached["Progress"]["B3"].value == 113
    assert cached["Progress"]["B7"].value == counts["Needs verification"]
    assert all(cached["Progress"][c.coordinate].value is not None
               for row in after["Progress"] for c in row if c.data_type == "f")
    candidate.replace(BOOK)
    print(f"Updated {len(UPDATES)} rows; all other workbook parts and formulas preserved.")
    print(dict(counts))


if __name__ == "__main__":
    main()
