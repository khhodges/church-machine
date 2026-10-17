"""One-time additive extension of the supplied CTMM workbook.

Refuses to overwrite an existing tracker or its manually maintained statuses.
Original sheet cells, formulas and styles are checked after serialization.
"""
from pathlib import Path
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "docs/phase-1/CTMM_W7_Threat_Test_Plan.xlsx"
PLAN = "docs/phase-1/W7-plus-plus-test-objectives.md"
HW = "docs/phase-1/04-release/isa-to-fpga-alignment-plan.md"
MINT = "docs/reviews/mint-phase1/open-issues.md"
ISA = "docs/phase-1/01-isa/instruction-set.md"


def snapshot(ws):
    return [(c.coordinate, c.value, str(c._style), c.number_format,
             c.hyperlink.target if c.hyperlink else None)
            for row in ws for c in row if c.value is not None]


def main():
    book = openpyxl.load_workbook(PATH)
    if "Completeness" in book.sheetnames:
        raise SystemExit("Tracker already exists. Edit it directly; do not regenerate statuses.")
    originals = {s.title: snapshot(s) for s in book}
    rows = []

    def add(id, area, requirement, status, finding, closure, source,
            scope="Phase 1 software", priority="High", dependencies=""):
        rows.append([id, area, requirement, scope, priority, status, "Unassigned",
                     dependencies, finding, closure, source, "2026-10-17", None, None])

    for row in book["Threat Map"].iter_rows(min_row=6, values_only=True):
        if isinstance(row[0], int):
            add(f"W7-{row[0]:02}", "W7 security", row[1],
                "Out of scope" if row[0] == 15 else "Not run",
                "Original workbook claim, not independently qualified. See original row.",
                str(row[6]) + " Coverage: " + str(row[7]),
                f"Workbook: Threat Map row {row[0]+5}; {PLAN}",
                "Phase 2" if row[0] == 15 else "Phase 1 assurance")
    for path, pattern, area, scope in [
        (HW, r"H\d\d", "Hardware acceptance", "Phase 1 hardware"),
        (HW, r"S\d\d", "Software acceptance", "Phase 1 software"),
        (PLAN, r"CM-[A-H]", "CM objectives", "Phase 1 assurance"),
        (PLAN, r"PERF-\d", "Performance", "Phase 1 measurement"),
        (MINT, r"G[1-8] .*", "Mint gates", "Phase 1 software"),
    ]:
        for line in (ROOT / path).read_text().splitlines():
            if not line.startswith("|"):
                continue
            cells = [x.strip().replace("**", "") for x in line.strip("|").split("|")]
            if not re.fullmatch(pattern, cells[0]):
                continue
            id = cells[0].split(" ")[0]
            requirement = cells[0] if area == "Mint gates" else cells[1]
            finding = ("Estimate only: " + " / ".join(cells[2:])) if area == "Performance" else (
                cells[1] if area == "Mint gates" else "Defined objective; full acceptance not established by this tracker.")
            add(id, area, requirement, "Open" if area == "Mint gates" else "Not run",
                finding, " / ".join(cells[2:]) if area != "Performance" else
                "Measure correct successful retirements for five one-second windows; retain exact build and trace mode.",
                path, scope, dependencies="See source gate dependencies")

    # Explicitly scoped work, not inferred approvals or invented test results.
    custom = [
        ("SPEC-01", "Specification", "Approved Master ISA and acceptance structure", "Complete",
         "Documentation exists; does not imply implementation conformance.",
         "Retain approved decisions and separate H01-H20 from S01-S07.", ISA),
        ("SPEC-02", "Specification", "Close remaining instruction oracles", "Open",
         "NV/reserved fields, shifts, TPERM, protected frames/Lambda and fault precedence remain explicit gaps.",
         "Approve independent expected words, flags, effects and faults before full conformance claims.", ISA),
        ("IDE-01", "IDE", "Additive Amobea and Thread design workspace", "Needs verification",
         "Integrated views exist; integrated acceptance is not re-run here.",
         "Verify navigation, retained drafts and no replacement of existing IDE tools.", "simulator/amobea/README.md"),
        ("IDE-02", "IDE", "Safe editor save, revision and buffer ownership", "Needs verification",
         "Workflow snapshot reports editor-roundtrip failures; root cause/current reproducibility not established.",
         "Reproduce isolated failures; verify create/edit/save, delayed responses, history and draft preservation.", "simulator/test_editor_roundtrip.js"),
        ("IDE-03", "IDE", "Namespace editor persistence and boot selection", "Needs verification",
         "Workflow snapshot reports ns-slot-modal-persist failures; no new diagnosis.",
         "Verify save/reload, resident selection and boot entry without silent layout/identity changes.", "server/lumps/ns-state.json"),
        ("IDE-04", "IDE", "Actionable diagnostics and executed-artifact freshness", "Needs verification",
         "Existing surfaces need integrated evidence, not generic failure alerts or inferred freshness.",
         "Show failed operation, actual reason, mutation state and corrective action; distinguish editor from executed bytes.", "simulator/test_execution_identity.js"),
        ("IDE-05", "IDE", "Hardware controls and causal command acknowledgement", "Needs verification",
         "Control/relay implementation exists; snapshot alone does not prove requested board action.",
         "Validate request-correlated acknowledgement, stale reconnect rejection, fault halt and simulator/board separation.", "tests/server/test_wukong_command_delivery.py"),
        ("IDE-06", "IDE", "Reference and compiler capability documentation agreement", "Needs verification",
         "Workflow snapshot flags API-reference and capability-block checks; not freshly reproduced.",
         "Reconcile approved ISA, compiler forms, reference search and examples; retain explicit historical distinctions.", ISA),
        ("FUNC-01", "Missing functions", "Authenticated original requester and live-use query", "Blocked",
         "No demonstrated legal service interface; protected frames do not grant software read authority.",
         "Approve and implement bounded attestation/traversal; no host callStack or widened heap workaround.", "docs/reviews/mint-phase1/context-feasibility.md"),
        ("FUNC-02", "Missing functions", "Same-scope Lambda entry and capability-result path", "Blocked",
         "G1 finds entry/continuation mismatch and unproven result authority.",
         "Resolve ISA contract; prove legal binding and result LOAD through real gates.", "docs/reviews/mint-phase1/context-feasibility.md"),
        ("FUNC-03", "Missing functions", "Tuple allocation and callable record lifecycle", "Needs verification",
         "58-word demonstration ABI and methods exist; not a new privacy mechanism.",
         "Verify allocation, bounds, permission revocation, malformed text and release/reuse with separate Threads for isolation.", "simulator/thread-heap-record.js"),
        ("FUNC-04", "Missing functions", "Dijkstra messages between isolated Threads", "Open",
         "Required communication direction; operational protocol/ABI is not demonstrated by the tuple lab.",
         "Define and approve capability-authorized send/receive, ownership, lifetime and fault semantics; test delivery.", "docs/phase-1/Amobea-V1.md"),
        ("FUNC-05", "Missing functions", "Continuation capture and one-shot checkpoint restore", "Blocked",
         "Design contract remains gated; external effects must not be replayed by heap rollback.",
         "Approve capture/restore scope, consumption identity and reset rules; test reuse/copy rejection.", "docs/phase-1/Amobea-V1.md"),
        ("FUNC-06", "Missing functions", "Protected Thread running ownership", "Blocked",
         "Live bit 13 remains proposed; saved-frame bit 13 is NIA, not ownership.",
         "Approve placement and all writer/validator changes; prove every activation path and self-target behavior.", "docs/phase-1/Amobea-V1.md"),
        ("FUNC-07", "Missing functions", "Reprogrammable startup breakout and application admission", "Open",
         "Startup diagnostics must allow open-ended checked handoff; not approval of live lazy loading.",
         "Specify target-selection authority and stopped resident delivery/activation; separate Phase 2 networking.", HW),
        ("SEAL-01", "Seal storage", "Define authoritative seal storage and lifetime", "Open",
         "Dedicated completeness item requested by Genesis; exact storage contract requires reconciliation, not an assumed missing implementation.",
         "Specify which seals, algorithm/meaning, trusted writer, location, access, retention and version binding; distinguish hashes, integrity32 and secret proofs.", "docs/reviews/mint-phase1/security.md"),
        ("SEAL-02", "Seal storage", "Bind seal to exact artifact and destination identity", "Needs verification",
         "Portable artifacts, localized C-lists and Namespace descriptors have different integrity boundaries.",
         "Reject tampering, stale/replayed seals and mismatched artifact/slot/generation; prove every covered field and permitted F/G mutation.", "docs/reviews/mint-phase1/namespace-update.md"),
        ("SEAL-03", "Seal storage", "Atomic seal/data publication and reset containment", "Open",
         "A successful checksum comparison does not prove exclusion or durable commit.",
         "Inject interruptions between data/seal writes; no partial or stale authority usable after reset; no implicit retry/rollback.", "docs/reviews/mint-phase1/namespace-update.md"),
        ("SEAL-04", "Seal storage", "Protected proof/key custody, backup and revocation", "Blocked",
         "G5/G6 open; public seals must not be confused with secret 128-bit proofs.",
         "Prove access isolation, confidential provisioning, atomic pool/ID/credential commit, anti-rollback and safe restoration.", "docs/reviews/mint-phase1/persistence.md"),
        ("SEAL-05", "Seal storage", "Storage fault detection versus authenticity", "Open",
         "No verified sideband SECDED coverage established here; integrity32 and UART CRC are not ECC.",
         "Record required storage fault coverage; any ECC design needs separate approval, preserves 32-bit GT and tests fault paths.", "hardware/hw_types.py"),
        ("SIM-01", "Simulator / ISA", "Independent current-ISA encoding and operand conformance", "Needs verification",
         "Assembler workflow reports failure; simulator agreement alone is not an independent oracle.",
         "Use hand-encoded vectors for all forms, bounds, conditions, retired encodings and CR13 restrictions.", ISA),
        ("SIM-02", "Simulator / ISA", "CALL RETURN LAMBDA cross-layer compatibility", "Blocked",
         "Synthetic component evidence reproduces a Lambda mismatch; green gap probes are not conformance.",
         "Same-byte simulator and RTL cases must match approved entry, continuation, frame and capability effects.", "docs/reviews/mint-phase1/context-feasibility.md"),
        ("SIM-03", "Simulator / ISA", "CHANGE SWITCH boot and Namespace geometry", "Needs verification",
         "Current fixtures and legacy saved image layouts differ; no silent binary reinterpretation.",
         "Compare current ISA and simulator/RTL using immutable fixture identities and exact initial state.", HW),
        ("SIM-04", "Simulator / ISA", "Watchdog and first-fault semantics", "Needs verification",
         "Elapsed board clocks cannot be inferred from browser step speed.",
         "Validate fault precedence and pre-recovery state; separate logical simulator checks from timed RTL/board proof.", HW),
        ("SIM-05", "Simulator / ISA", "Historical artifact adaptation provenance", "Needs verification",
         "Alice/Mallory fixtures adapt old code in isolated RAM; originals remain immutable.",
         "Record exact original/adapted words and current ISA; never claim adapted runs qualify historical binary unchanged.", "simulator/test_alice_mallory_lab.js"),
        ("REL-01", "RTL / bitstream", "Fresh exact-source build provenance", "Open",
         "Prior inspection found 22/39 recorded build inputs differ; this is not a new build verification.",
         "Regenerate authorized exact inputs; verify hashes, ROM/configuration, timing, DRC and artifact sidecars.", "docs/wukong-release-chain.md"),
        ("REL-02", "RTL / bitstream", "Physical acceptance of exact release candidate", "Not run",
         "No physical Phase 1 acceptance performed for this tracker.",
         "After separate authorization, run H19/H20 with exact board/build identity and retained raw observations.", HW),
        ("REL-03", "RTL / bitstream", "Resolve reported regression failures", "Needs verification",
         "Snapshot reports LUMP consistency, boot image, endpoints, ROM guards and E2E failures; causes unverified.",
         "Isolate current failures, classify stale fixtures versus defects; record commands, output, identities and reruns.", "scripts/run-all-tests.sh"),
    ]
    for id, area, req, status, finding, closure, source in custom:
        add(id, area, req, status, finding, closure, source,
            "Phase 1 hardware" if area == "RTL / bitstream" else "Phase 1 assurance")
    for name, role in [
        ("Boot.NS", "Namespace foundation"), ("Boot.Thread", "Initial context"),
        ("Thread.2", "Additional context"), ("Thread.3", "Additional context"),
        ("SelfTest", "Startup diagnostic"), ("CapabilityTest", "Capability/cooperative test coordinator"),
        ("WukongCallHome", "Authorized UART reporting"), ("ide.Mint", "Object and credential service"),
        ("ide.PkVerify", "Exact-request PassKey verifier"), ("ide.NSTab", "Namespace table RW object"),
        ("ide.Alice", "Authorized containment example"), ("ide.Mallory", "Adversarial containment example"),
    ]:
        add("RES-" + name, "Named abstractions", name + " — " + role,
            "Open" if name in ("ide.Mint", "ide.PkVerify", "ide.NSTab") else "Needs verification",
            "Saved selection/library availability is not physical residency or current-ISA qualification.",
            "Record exact artifact, approved destination/binding, resident footprint and simulator/RTL/board acceptance separately.",
            "server/lumps/ns-state.json; " + (MINT if name.startswith("ide.") else HW))

    assert len({r[0] for r in rows}) == len(rows)
    named_findings = {
        "Boot.NS": "Saved snapshot uses legacy slot-2 UART layout; approved Phase 1 entry layout must be reconciled before qualification.",
        "Boot.Thread": "Saved resident initial context; not evidence of exact-board Phase 1 acceptance.",
        "Thread.2": "Saved resident context; approved entry and cooperative handoff need exact-artifact qualification.",
        "Thread.3": "Saved resident context; approved entry and cooperative handoff need exact-artifact qualification.",
        "SelfTest": "Saved resident startup diagnostic; discarded failure results can mask errors. Inspect branch and return-status evidence.",
        "CapabilityTest": "Saved resident boot target/test coordinator; validate calls to SelfTest and WukongCallHome, returns and cooperative CHANGE.",
        "WukongCallHome": "Saved resident UART reporting routine. Receipt alone does not prove all tests passed, current bitstream or authenticated identity.",
        "ide.Mint": "57-answer ledger and Genesis review established; G1-G8 open. First credential Create-only in Boot, with permanently empty Replace/Destroy targets. Absent from inspected saved Namespace.",
        "ide.PkVerify": "Designed exact-request verifier, not qualified operational service. Absent from inspected saved Namespace; protected context and provisioning still gated.",
        "ide.NSTab": "Specified RW table object, not an executable helper. Only Mint's internal Lambda may hold its table access; absent from inspected saved Namespace.",
        "ide.Alice": "Merged isolated simulator fixtures cover own-tuple DWRITE/DREAD and adapted legacy scenarios; not a new run or board acceptance. Not a saved resident here.",
        "ide.Mallory": "Merged seven tuple scenarios include absent-capability denial, own-tuple positive control, foreign/range read/write denial and full 58-word Alice integrity. No new run or board proof.",
    }
    for row in rows:
        if row[0].startswith("RES-"):
            name = row[0][4:]
            row[8] = named_findings[name]
            if name in ("ide.Alice", "ide.Mallory"):
                row[10] = "simulator/test_alice_mallory_lab.js; simulator/test_alice_mallory_tuple_lab.js"
        if row[0] in ("FUNC-01", "RES-ide.PkVerify"):
            row[7] = "G2; G3; G6"
        if row[0] == "FUNC-02":
            row[7] = "G1; SPEC-02"
        if row[0] == "SEAL-04":
            row[7] = "G5; G6"
        if row[0] == "REL-02":
            row[7] = "REL-01; H19; H20; separate hardware authorization"
        if row[0] == "RES-ide.Mint":
            row[7] = "G1-G8"
    for row in rows:
        for source in row[10].split("; "):
            if not source.startswith("Workbook:"):
                assert (ROOT / source).is_file(), source
    ws = book.create_sheet("Completeness")
    ws.append(["ID", "Area", "Requirement / issue", "Scope", "Priority", "Status",
               "Owner", "Dependencies", "Current evidence / gap", "Closure criterion / next action",
               "Source reference", "Reviewed date", "Evidence / result link", "Due date"])
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "D2"
    tab = Table(displayName="Phase1Tracker", ref=f"A1:N{ws.max_row}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    statuses = ["Open", "Blocked", "Not run", "Needs verification", "In progress",
                "Complete", "Failed", "Out of scope"]
    for col, values in [("F", statuses), ("E", ["Critical", "High", "Medium", "Low"])]:
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"')
        dv.errorTitle, dv.error = "Choose a listed value", "Use the dropdown to preserve summary counts."
        dv.showErrorMessage, dv.errorStyle = True, "stop"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}10000")
    for status, color in [("Complete", "D9EAD3"), ("Blocked", "F4CCCC"),
                          ("Failed", "F4CCCC"), ("Needs verification", "FFF2CC")]:
        ws.conditional_formatting.add("F2:F10000", FormulaRule(
            formula=[f'F2="{status}"'], fill=PatternFill("solid", fgColor=color)))
    dashboard = book.create_sheet("Progress")
    dashboard.append(["Phase 1 completeness — evidence-based tracking", "Count / rate"])
    dashboard.append(["Updated", "2026-10-17"])
    dashboard.append(["Total tracked rows", '=COUNTA(Completeness!A2:A10000)'])
    for status in statuses:
        dashboard.append([status, f'=COUNTIF(Completeness!F2:F10000,"{status}")'])
    dashboard.append(["In-scope completion", '=IFERROR(COUNTIF(Completeness!F2:F10000,"Complete")/(COUNTA(Completeness!A2:A10000)-COUNTIF(Completeness!F2:F10000,"Out of scope")),0)'])
    dashboard.cell(dashboard.max_row, 2).number_format = "0.0%"
    dashboard.append(["Interpretation", "Unweighted row closure, NOT percentage of ISA/security proved. Software gaps are not automatically bitstream gates."])
    dashboard.append(["Area", "Tracked", "Complete", "Blocked"])
    for area in sorted({r[1] for r in rows}):
        n = dashboard.max_row + 1
        dashboard.append([area, f'=COUNTIF(Completeness!B2:B10000,A{n})',
                          f'=COUNTIFS(Completeness!B2:B10000,A{n},Completeness!F2:F10000,"Complete")',
                          f'=COUNTIFS(Completeness!B2:B10000,A{n},Completeness!F2:F10000,"Blocked")'])
    guide = book.create_sheet("Tracking guide")
    for r in [
        ("Rule", "Meaning"),
        ("Purpose", "Expanded CTMM W7 completeness tracker; original Threat Map, Summary and Method retained without changed cells/formulas/styles."),
        ("Original claims", "Original Ratified/Proven labels are supplied claims, not newly verified findings. Original Summary covers only the 15 threats; Progress covers the new register."),
        ("Complete", "Closure criterion satisfied at its stated evidence level; cite dated reproducible evidence. Documentation completion is not hardware acceptance."),
        ("Needs verification", "Implementation, reported merge or failure exists, but current end-to-end acceptance/root cause has not been verified here."),
        ("Open / Blocked", "Open: work or decision remains. Blocked: prerequisite authority, contract or mechanism missing. Not run: defined objective with no qualifying run recorded."),
        ("No invented status", "No new simulator, RTL or hardware runs were performed to populate this workbook; workflow snapshot failures are leads, not confirmed current defects."),
        ("Maintaining progress", "Edit Completeness, assign Owner and Due date, add evidence, update Reviewed date and dependencies. Append rows within the Excel table; summaries cover rows 2–10000."),
        ("Evidence identity", "Record exact source, fixture/artifact, build/board, command, date, outcome and exclusions. Never put credentials, keys, proofs or private records in this workbook."),
        ("Scope", "Phase 1 hardware and software delivery are separate gates. Phase 2 networking remains out of scope. Proposed ECC and running-bit interfaces are not approved by this tracker."),
        ("Seal storage", "Reconcile authoritative seal semantics/storage with existing implementations; distinguish authenticity seals, public integrity checks, secret proof custody and ECC."),
        ("Calculations", "Excel/LibreOffice recalculates formulas on open. Python does not evaluate them. Completion percentage is unweighted row closure, not a security-confidence metric."),
        ("Source references", "Repository-relative source paths describe where to find requirements; they are not portable public URLs or runtime test evidence."),
    ]:
        guide.append(r)
    for sheet in [ws, dashboard, guide]:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = sheet.freeze_panes or "A2"
        for c in sheet[1]:
            c.fill = PatternFill("solid", fgColor="17365D")
            c.font = Font(color="FFFFFF", bold=True)
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    for col, width in zip("ABCDEFGHIJKLMN", [23, 23, 48, 24, 12, 23, 20, 28, 65, 70, 60, 16, 50, 16]):
        ws.column_dimensions[col].width = width
    for n in range(2, ws.max_row + 1):
        ws.row_dimensions[n].height = 90
    for sheet in [dashboard, guide]:
        sheet.column_dimensions["A"].width = 52
        sheet.column_dimensions["B"].width = 95 if sheet == guide else 48
        sheet.column_dimensions["C"].width = 16
        sheet.column_dimensions["D"].width = 16
        for n in range(2, sheet.max_row + 1):
            sheet.row_dimensions[n].height = 48
    book.active = book.index(dashboard)
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    candidate = PATH.with_suffix(".candidate.xlsx")
    book.save(candidate)
    check = openpyxl.load_workbook(candidate)
    assert all(snapshot(check[name]) == data for name, data in originals.items())
    assert check["Completeness"].max_row == len(rows) + 1
    assert len(check["Completeness"].tables) == 1
    assert len(check["Completeness"].data_validations.dataValidation) == 2
    candidate.replace(PATH)
    print(f"Validated {len(rows)} tracker rows; all three original sheet cells/formulas/styles unchanged.")
    print("Sheets:", ", ".join(check.sheetnames))


if __name__ == "__main__":
    main()
